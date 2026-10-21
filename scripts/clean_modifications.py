import os
import shutil
import datetime
import openpyxl
import pandas as pd

def backup_and_clean_modifications():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target_file = os.path.join(base_dir, "MODIFICACIONES_SABADOS.xlsx")
    
    if not os.path.exists(target_file):
        print(f"Error: No se encontró el archivo {target_file}")
        return False
        
    # 1. Crear respaldo fechado
    timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = os.path.join(base_dir, f"MODIFICACIONES_SABADOS_backup_{timestamp_str}.xlsx")
    shutil.copy2(target_file, backup_file)
    print(f"[OK] Respaldo creado exitosamente en: {backup_file}")
    
    # 2. Cargar datos actuales
    wb_orig = openpyxl.load_workbook(target_file)
    sheet_names = wb_orig.sheetnames
    print(f"Hojas encontradas: {sheet_names}")
    
    # 3. Procesar hoja MODIFICACIONES
    df_mods = pd.read_excel(target_file, sheet_name="MODIFICACIONES")
    total_original = len(df_mods)
    counts_orig = df_mods['TYPE'].value_counts().to_dict()
    print(f"Total registros originales en MODIFICACIONES: {total_original}")
    print(f"Desglose original por TYPE: {counts_orig}")
    
    # Filtrar removiendo estrictamente ELIMINAR
    df_clean = df_mods[df_mods['TYPE'] != 'ELIMINAR'].copy()
    total_cleaned = len(df_clean)
    counts_clean = df_clean['TYPE'].value_counts().to_dict()
    print(f"Total registros depurados (sin ELIMINAR): {total_cleaned}")
    print(f"Desglose limpio por TYPE: {counts_clean}")
    
    # Reasignar IDs correlativos (1, 2, 3...)
    df_clean['ID'] = range(1, total_cleaned + 1)
    
    # 4. Crear nuevo libro limpio conservando todas las hojas
    new_wb = openpyxl.Workbook()
    # Eliminar hoja default
    default_sheet = new_wb.active
    
    # Escribir hoja MODIFICACIONES
    ws_mods = new_wb.create_sheet(title="MODIFICACIONES")
    headers = list(df_clean.columns)
    ws_mods.append(headers)
    for row in df_clean.itertuples(index=False):
        # Convertir timestamps o NaNs apropiadamente
        clean_row = []
        for val in row:
            if pd.isna(val):
                clean_row.append(None)
            elif isinstance(val, (datetime.date, datetime.datetime)):
                clean_row.append(val.strftime("%Y-%m-%d") if isinstance(val, datetime.date) and not isinstance(val, datetime.datetime) else val.strftime("%Y-%m-%d %H:%M:%S"))
            else:
                clean_row.append(val)
        ws_mods.append(clean_row)
        
    # Copiar otras hojas (ej. PERSONAL_MANUAL si existe)
    for s_name in sheet_names:
        if s_name != "MODIFICACIONES":
            df_other = pd.read_excel(target_file, sheet_name=s_name)
            ws_other = new_wb.create_sheet(title=s_name)
            ws_other.append(list(df_other.columns))
            for row in df_other.itertuples(index=False):
                clean_row = [None if pd.isna(v) else v for v in row]
                ws_other.append(clean_row)
            print(f"Copiada hoja adicional: {s_name} ({len(df_other)} filas)")
            
    new_wb.remove(default_sheet)
    new_wb.save(target_file)
    new_wb.close()
    wb_orig.close()
    
    print(f"[OK] Archivo {target_file} depurado y guardado exitosamente.")
    return True

if __name__ == "__main__":
    backup_and_clean_modifications()
